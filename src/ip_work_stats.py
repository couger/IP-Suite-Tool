#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""工作统计GUI界面"""

import os
import re
import time
import threading
import glob
from dataclasses import dataclass
from datetime import datetime
from typing import Iterable, Optional

import tkinter as tk
from tkinter import ttk, messagebox, filedialog

from openpyxl import Workbook, load_workbook

try:
    import xlrd
except ImportError:
    xlrd = None


EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls"}
TXT_EXTENSIONS = {".txt"}
WORD_EXTENSIONS = {".docx"}
SUPPORTED_EXTENSIONS = EXCEL_EXTENSIONS | TXT_EXTENSIONS | WORD_EXTENSIONS
PREFIX_RE = re.compile(r"^(pl\d{12}|add\d{12})(.*)$", re.IGNORECASE)
HEADER_SCAN_ROWS = 10


@dataclass
class WorkStatRecord:
    """单个工作统计结果"""

    sequence: int
    prefix: str
    title: str
    applicant_unit: str
    applicant: str
    source_file: str
    status: str = "成功"
    message: str = ""


class WorkStats:
    """工作统计页面"""

    def __init__(self, parent):
        self.parent = parent
        self.selected_files = []
        self.is_processing = False
        self.processing_thread = None
        self.embed_attachments_var = tk.BooleanVar(value=False)

        self._init_ui()

        self.parent.after(
            100,
            lambda: self.import_files_btn.focus_set()
            if hasattr(self, "import_files_btn")
            else None,
        )

    def _init_ui(self):
        self.main_frame = ttk.Frame(self.parent)
        self.main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        self._create_import_section()
        self._create_output_section()
        self._create_processing_section()
        self._create_results_section()

    def _create_import_section(self):
        """创建与IP脱敏器一致的文件导入区域"""
        import_frame = ttk.LabelFrame(self.main_frame, text="文件导入", padding="6")
        import_frame.pack(fill=tk.X, pady=(0, 6))

        main_container = ttk.Frame(import_frame)
        main_container.pack(fill=tk.BOTH, expand=True)

        list_frame = ttk.Frame(main_container)
        list_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        main_container.columnconfigure(0, weight=1)
        main_container.rowconfigure(0, weight=1)

        columns = ("文件名", "路径", "大小", "类型")
        self.files_tree = ttk.Treeview(
            list_frame,
            columns=columns,
            show="headings",
            height=4,
            selectmode="extended",
        )

        self.files_tree.heading("文件名", text="文件名")
        self.files_tree.heading("路径", text="路径")
        self.files_tree.heading("大小", text="大小")
        self.files_tree.heading("类型", text="类型")

        self.files_tree.column("文件名", width=160)
        self.files_tree.column("路径", width=220)
        self.files_tree.column("大小", width=70)
        self.files_tree.column("类型", width=70)

        scrollbar = ttk.Scrollbar(
            list_frame, orient=tk.VERTICAL, command=self.files_tree.yview
        )
        self.files_tree.configure(yscrollcommand=scrollbar.set)

        self.files_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        button_frame = ttk.Frame(main_container)
        button_frame.grid(row=0, column=1, sticky="ns")
        button_frame.columnconfigure(0, weight=1)

        button_width = 12
        self.import_files_btn = ttk.Button(
            button_frame,
            text="导入文件",
            command=self._import_files,
            width=button_width,
        )
        self.import_files_btn.grid(row=0, column=0, pady=(0, 6), sticky="ew")

        self.import_folder_btn = ttk.Button(
            button_frame,
            text="导入目录",
            command=self._import_folder,
            width=button_width,
        )
        self.import_folder_btn.grid(row=1, column=0, pady=(0, 6), sticky="ew")

        self.delete_files_btn = ttk.Button(
            button_frame,
            text="删除选中",
            command=self._delete_selected_files,
            width=button_width,
        )
        self.delete_files_btn.grid(row=2, column=0, pady=(0, 6), sticky="ew")

        self.clear_files_btn = ttk.Button(
            button_frame,
            text="清空列表",
            command=self._clear_files_list,
            width=button_width,
        )
        self.clear_files_btn.grid(row=3, column=0, pady=(0, 6), sticky="ew")

    def _create_output_section(self):
        output_frame = ttk.LabelFrame(self.main_frame, text="输出设置", padding="6")
        output_frame.pack(fill=tk.X, pady=(0, 6))

        ttk.Label(output_frame, text="输出文件:").pack(side=tk.LEFT)
        self.output_path_var = tk.StringVar()
        self.output_entry = ttk.Entry(
            output_frame, textvariable=self.output_path_var, width=40
        )
        self.output_entry.pack(side=tk.LEFT, padx=(10, 6), fill=tk.X, expand=True)

        self.browse_output_btn = ttk.Button(
            output_frame, text="浏览", command=self._browse_output_file
        )
        self.browse_output_btn.pack(side=tk.LEFT)

        self.embed_attachments_check = ttk.Checkbutton(
            output_frame,
            text="插入文档附件",
            variable=self.embed_attachments_var,
            command=self._on_embed_attachments_toggle,
        )
        self.embed_attachments_check.pack(side=tk.LEFT, padx=(10, 0))

    def _create_processing_section(self):
        processing_frame = ttk.LabelFrame(self.main_frame, text="统计控制", padding="6")
        processing_frame.pack(fill=tk.X, pady=(0, 6))

        main_container = ttk.Frame(processing_frame)
        main_container.pack(fill=tk.X, expand=True)

        progress_frame = ttk.Frame(main_container)
        progress_frame.grid(row=0, column=0, sticky="ew", padx=(0, 10))
        main_container.columnconfigure(0, weight=1)

        self.progress_var = tk.DoubleVar()
        self.progress_bar = ttk.Progressbar(
            progress_frame,
            variable=self.progress_var,
            maximum=100,
            mode="determinate",
            style="Green.Horizontal.TProgressbar",
        )
        self.progress_bar.pack(fill=tk.X, pady=(0, 4))

        self.status_var = tk.StringVar(value="就绪")
        self.status_label = ttk.Label(progress_frame, textvariable=self.status_var)
        self.status_label.pack(anchor=tk.W)

        button_frame = ttk.Frame(main_container)
        button_frame.grid(row=0, column=1, sticky="ns")
        button_frame.columnconfigure(0, weight=1)

        self.stats_btn = ttk.Button(
            button_frame,
            text="工作统计",
            command=self._start_statistics,
            style="Accent.TButton",
            width=12,
        )
        self.stats_btn.grid(row=0, column=0, pady=(0, 6), sticky="ew")

    def _create_results_section(self):
        results_frame = ttk.LabelFrame(self.main_frame, text="统计结果", padding="6")
        results_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 6))

        columns = ("序号", "前缀", "名称", "组织", "联系人", "状态")
        self.results_tree = ttk.Treeview(
            results_frame, columns=columns, show="headings", height=8
        )

        for column in columns:
            self.results_tree.heading(column, text=column)

        self.results_tree.column("序号", width=50, anchor=tk.CENTER)
        self.results_tree.column("前缀", width=130)
        self.results_tree.column("名称", width=280)
        self.results_tree.column("组织", width=180)
        self.results_tree.column("联系人", width=120)
        self.results_tree.column("状态", width=80)

        scrollbar = ttk.Scrollbar(
            results_frame, orient=tk.VERTICAL, command=self.results_tree.yview
        )
        self.results_tree.configure(yscrollcommand=scrollbar.set)

        self.results_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

    def _format_file_size(self, size_bytes: int) -> str:
        if size_bytes < 1024:
            return f"{size_bytes} B"
        if size_bytes < 1024 * 1024:
            return f"{size_bytes / 1024:.1f} KB"
        return f"{size_bytes / (1024 * 1024):.1f} MB"

    def _file_type_label(self, file_path: str) -> str:
        ext = os.path.splitext(file_path)[1].lower()
        if ext in EXCEL_EXTENSIONS:
            return "Excel文件"
        if ext in TXT_EXTENSIONS:
            return "文本文件"
        if ext in WORD_EXTENSIONS:
            return "Word文件"
        return "未知"

    def _is_excel_file(self, file_path: str) -> bool:
        return os.path.splitext(file_path)[1].lower() in EXCEL_EXTENSIONS

    def _is_supported_file(self, file_path: str) -> bool:
        file_name = os.path.basename(file_path)
        if file_name.startswith("~$"):
            return False
        return os.path.splitext(file_path)[1].lower() in SUPPORTED_EXTENSIONS

    def _set_controls_state(self, state):
        self.import_files_btn.config(state=state)
        self.import_folder_btn.config(state=state)
        self.delete_files_btn.config(state=state)
        self.clear_files_btn.config(state=state)
        self.browse_output_btn.config(state=state)
        self.embed_attachments_check.config(state=state)
        self.stats_btn.config(state=state)

    def _on_embed_attachments_toggle(self):
        if not self.embed_attachments_var.get():
            return

        messagebox.showwarning(
            "附件插入前提条件",
            "插入文档附件需要满足以下条件：\n\n"
            "1. 当前电脑已安装 Microsoft Excel；\n"
            "2. Excel 允许通过“插入对象”功能嵌入文件；\n"
            "3. Office 安全策略未禁用 OLE 对象插入；\n"
            "4. 统计时请不要手动操作正在自动化处理的 Excel 窗口。\n\n"
            "该功能会明显增加统计耗时，文件数量较多时请耐心等待。",
        )

    def _add_files_to_list(self, file_paths: Iterable[str], show_message: bool = True):
        added_count = 0
        skipped_count = 0
        for file_path in file_paths:
            file_path = os.path.normpath(file_path)
            if not os.path.isfile(file_path) or file_path in self.selected_files:
                continue
            if not self._is_supported_file(file_path):
                skipped_count += 1
                continue

            self.selected_files.append(file_path)
            file_name = os.path.basename(file_path)
            file_dir = os.path.dirname(file_path)
            try:
                file_size = self._format_file_size(os.path.getsize(file_path))
            except (OSError, FileNotFoundError):
                file_size = "无法获取"

            self.files_tree.insert(
                "",
                tk.END,
                values=(file_name, file_dir, file_size, self._file_type_label(file_path)),
            )
            added_count += 1

        if added_count and not self.output_path_var.get():
            self.output_path_var.set(self._get_default_output_file())

        if added_count > 0 and show_message:
            message = f"成功导入 {added_count} 个文件"
            if skipped_count:
                message += f"\n已跳过 {skipped_count} 个非支持类型文件"
            messagebox.showinfo("导入完成", message)
        elif skipped_count and show_message:
            messagebox.showwarning(
                "未导入文件",
                "未找到可导入文件。\n支持类型：Excel、TXT、Word(docx)。",
            )

        return added_count

    def load_files(self, file_paths: Iterable[str]):
        """从其他标签页载入文件列表"""
        self._clear_files_list(show_message=False)
        added_count = self._add_files_to_list(file_paths, show_message=False)
        self.status_var.set(f"已从IP脱敏器载入 {added_count} 个文件")

    def _import_files(self):
        if self.is_processing:
            messagebox.showwarning("处理中", "请等待当前统计完成后再导入文件")
            return

        file_paths = filedialog.askopenfilenames(
            title="选择要导入的文件",
            filetypes=[
                ("支持文件", "*.xlsx *.xlsm *.xltx *.xltm *.xls *.txt *.docx"),
                ("Excel文件", "*.xlsx *.xlsm *.xltx *.xltm *.xls"),
                ("文本文件", "*.txt"),
                ("Word文件", "*.docx"),
            ],
        )
        if file_paths:
            self._add_files_to_list(file_paths)

    def _import_folder(self):
        if self.is_processing:
            messagebox.showwarning("处理中", "请等待当前统计完成后再导入文件")
            return

        folder_path = filedialog.askdirectory(title="选择要导入的目录")
        if not folder_path:
            return

        file_paths = []
        for root, _, files in os.walk(folder_path):
            for file_name in files:
                file_path = os.path.join(root, file_name)
                if self._is_supported_file(file_path):
                    file_paths.append(file_path)

        if file_paths:
            self._add_files_to_list(file_paths)
        else:
            messagebox.showwarning(
                "未导入文件",
                "目录中没有可导入文件。\n支持类型：Excel、TXT、Word(docx)。",
            )

    def _delete_selected_files(self):
        if self.is_processing:
            messagebox.showwarning("处理中", "请等待当前统计完成后再删除文件")
            return

        selected_items = self.files_tree.selection()
        if not selected_items:
            messagebox.showwarning("无选择", "请先选择要删除的文件")
            return

        for item in selected_items:
            values = self.files_tree.item(item, "values")
            if values and len(values) >= 2:
                file_path = os.path.normpath(os.path.join(values[1], values[0]))
                if file_path in self.selected_files:
                    self.selected_files.remove(file_path)
            self.files_tree.delete(item)

    def _clear_files_list(self, show_message: bool = True):
        if self.is_processing:
            messagebox.showwarning("处理中", "请等待当前统计完成后再清空列表")
            return

        self.selected_files.clear()
        self.files_tree.delete(*self.files_tree.get_children())
        self.results_tree.delete(*self.results_tree.get_children())
        self.output_path_var.set("")
        self.progress_var.set(0)
        self.status_var.set("就绪")

        if show_message:
            messagebox.showinfo("清空完成", "文件列表已清空")

    def _browse_output_file(self):
        file_path = filedialog.asksaveasfilename(
            title="保存工作统计结果",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx")],
            initialfile=f"工作统计_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        )
        if file_path:
            self.output_path_var.set(os.path.normpath(file_path))

    def _get_default_output_file(self) -> str:
        if self.selected_files:
            output_dir = os.path.dirname(self.selected_files[-1])
        else:
            output_dir = os.getcwd()
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return os.path.join(output_dir, f"工作统计_{timestamp}.xlsx")

    def _start_statistics(self):
        if self.is_processing:
            return

        if not self.selected_files:
            messagebox.showwarning("无文件", "请先导入要统计的文件")
            return

        target_files = [path for path in self.selected_files if self._is_supported_file(path)]
        if not target_files:
            messagebox.showwarning(
                "无可统计文件",
                "当前列表中没有可统计文件。\n支持类型：Excel、TXT、Word(docx)。",
            )
            return

        if not self.output_path_var.get():
            self.output_path_var.set(self._get_default_output_file())

        self.is_processing = True
        self._set_controls_state(tk.DISABLED)
        self.results_tree.delete(*self.results_tree.get_children())
        self.progress_var.set(0)
        self.status_var.set("正在统计...")

        output_path = self.output_path_var.get()
        embed_attachments = self.embed_attachments_var.get()
        self.processing_thread = threading.Thread(
            target=self._statistics_thread,
            args=(target_files, output_path, embed_attachments),
            daemon=True,
        )
        self.processing_thread.start()

    def _statistics_thread(self, target_files: list[str], output_path: str, embed_attachments: bool):
        try:
            records = self._collect_records(target_files)
            self._save_records(
                records,
                output_path,
                embed_attachments=embed_attachments,
            )

            self.parent.after(0, lambda: self._show_records(records))
            self.parent.after(
                0,
                lambda: messagebox.showinfo(
                    "统计完成",
                    f"已统计 {len(records)} 个目标文件。\n结果已保存到:\n{output_path}",
                ),
            )
        except Exception as exc:
            error_message = str(exc)
            self.parent.after(
                0, lambda message=error_message: messagebox.showerror("统计错误", message)
            )
        finally:
            self.parent.after(0, self._reset_ui_state)

    def _collect_records(self, target_files: Optional[list[str]] = None):
        records = []
        if target_files is None:
            target_files = [
                path for path in self.selected_files if self._is_supported_file(path)
            ]
        total = len(target_files)

        for index, file_path in enumerate(target_files, 1):
            record = self._read_record(index, file_path)
            records.append(record)

            progress = (index / total) * 100 if total else 100
            self.parent.after(0, lambda value=progress: self.progress_var.set(value))
            self.parent.after(
                0,
                lambda current=index, count=total: self.status_var.set(
                    f"正在统计... ({current}/{count})"
                ),
            )

        return records

    def _read_record(self, sequence: int, file_path: str) -> WorkStatRecord:
        stem = os.path.splitext(os.path.basename(file_path))[0]
        match = PREFIX_RE.match(stem)
        prefix = match.group(1) if match else ""
        title = match.group(2).lstrip("-_ ") if match else stem

        if not self._is_excel_file(file_path):
            return WorkStatRecord(
                sequence=sequence,
                prefix=prefix,
                title=title,
                applicant_unit="",
                applicant="",
                source_file=file_path,
                message="TXT/Word文件不读取组织和联系人",
            )

        try:
            applicant_unit, applicant = self._read_excel_applicant_info(file_path)
            return WorkStatRecord(
                sequence=sequence,
                prefix=prefix,
                title=title,
                applicant_unit=applicant_unit,
                applicant=applicant,
                source_file=file_path,
            )
        except Exception as exc:
            return WorkStatRecord(
                sequence=sequence,
                prefix=prefix,
                title=title,
                applicant_unit="",
                applicant="",
                source_file=file_path,
                status="失败",
                message=str(exc),
            )

    def _read_excel_applicant_info(self, file_path: str) -> tuple[str, str]:
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".xls":
            return self._read_xls_applicant_info(file_path)
        return self._read_openpyxl_applicant_info(file_path)

    def _read_openpyxl_applicant_info(self, file_path: str) -> tuple[str, str]:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            header_row_index = None
            unit_col = None
            applicant_col = None

            for row_index, row in enumerate(
                sheet.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS), 1
            ):
                headers = [
                    "" if cell.value is None else str(cell.value).strip()
                    for cell in row
                ]
                try:
                    unit_col = self._find_header_index(headers, "组织")
                    applicant_col = self._find_header_index(headers, "联系人")
                    header_row_index = row_index
                    break
                except ValueError:
                    continue

            if header_row_index is None or unit_col is None or applicant_col is None:
                raise ValueError("前10行未找到包含“组织”和“联系人”的标题行")

            applicant_unit = ""
            applicant = ""
            for row in sheet.iter_rows(min_row=header_row_index + 1, values_only=True):
                if unit_col < len(row) and row[unit_col] not in (None, ""):
                    applicant_unit = str(row[unit_col]).strip()
                if applicant_col < len(row) and row[applicant_col] not in (None, ""):
                    applicant = str(row[applicant_col]).strip()

            return applicant_unit, applicant
        finally:
            workbook.close()

    def _read_xls_applicant_info(self, file_path: str) -> tuple[str, str]:
        if xlrd is None:
            raise RuntimeError("当前环境未安装xlrd，无法读取xls文件")

        workbook = xlrd.open_workbook(file_path)
        sheet = workbook.sheet_by_index(0)
        if sheet.nrows == 0:
            raise ValueError("Excel文件为空")

        header_row_index = None
        unit_col = None
        applicant_col = None
        for row_idx in range(min(sheet.nrows, HEADER_SCAN_ROWS)):
            headers = [
                str(sheet.cell_value(row_idx, col)).strip()
                for col in range(sheet.ncols)
            ]
            try:
                unit_col = self._find_header_index(headers, "组织")
                applicant_col = self._find_header_index(headers, "联系人")
                header_row_index = row_idx
                break
            except ValueError:
                continue

        if header_row_index is None or unit_col is None or applicant_col is None:
            raise ValueError("前10行未找到包含“组织”和“联系人”的标题行")

        applicant_unit = ""
        applicant = ""
        for row_idx in range(header_row_index + 1, sheet.nrows):
            unit_value = sheet.cell_value(row_idx, unit_col)
            applicant_value = sheet.cell_value(row_idx, applicant_col)
            if unit_value not in (None, ""):
                applicant_unit = str(unit_value).strip()
            if applicant_value not in (None, ""):
                applicant = str(applicant_value).strip()

        return applicant_unit, applicant

    def _find_header_index(self, headers: list[str], keyword: str) -> int:
        aliases = {"组织": ("组织", "申请单位"), "联系人": ("联系人", "申请人")}
        names = aliases.get(keyword, (keyword,))
        exact = [index for index, value in enumerate(headers) if value.strip() == keyword]
        matches = exact or [index for index, value in enumerate(headers)
                            if any(name in value for name in names)]
        if not matches:
            raise ValueError(f"标题行未找到包含“{keyword}”的列")
        return matches[-1]

    def _save_records(
        self,
        records: list[WorkStatRecord],
        output_path: str,
        embed_attachments: bool = False,
    ):
        output_dir = os.path.dirname(output_path)
        if output_dir:
            os.makedirs(output_dir, exist_ok=True)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "工作统计"

        headers = [
            "序号",
            "前缀",
            "文件名剩余部分",
            "组织",
            "联系人",
            "附件",
            "源文件路径",
            "状态",
            "备注",
        ]
        sheet.append(headers)

        for record in records:
            sheet.append(
                [
                    record.sequence,
                    record.prefix,
                    record.title,
                    record.applicant_unit,
                    record.applicant,
                    "",
                    record.source_file,
                    record.status,
                    record.message,
                ]
            )

        widths = [8, 18, 46, 28, 18, 34, 60, 12, 38]
        for index, width in enumerate(widths, 1):
            column_letter = sheet.cell(row=1, column=index).column_letter
            sheet.column_dimensions[column_letter].width = width
        for row_index in range(2, len(records) + 2):
            sheet.row_dimensions[row_index].height = 45

        workbook.save(output_path)

        if embed_attachments:
            self._embed_attachments_as_icons(output_path, records)

    def _embed_attachments_as_icons(
        self, output_path: str, records: list[WorkStatRecord]
    ) -> None:
        """使用 Windows Excel COM 将源文件作为OLE图标附件插入统计表。"""
        try:
            import pythoncom
            import win32com.client
        except ImportError as exc:
            raise RuntimeError("当前环境未安装pywin32，无法插入OLE附件图标") from exc

        pythoncom.CoInitialize()
        excel = None
        workbook = None

        def retry_excel_call(action, description: str):
            last_error = None
            for attempt in range(12):
                try:
                    return action()
                except Exception as exc:
                    last_error = exc
                    pythoncom.PumpWaitingMessages()
                    time.sleep(0.4 + attempt * 0.15)
            raise RuntimeError(f"{description}失败") from last_error

        try:
            excel = win32com.client.gencache.EnsureDispatch("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False

            workbook = retry_excel_call(
                lambda: excel.Workbooks.Open(os.path.abspath(output_path)),
                "打开工作统计Excel文件",
            )
            worksheet = retry_excel_call(
                lambda: workbook.Worksheets("工作统计"),
                "定位工作统计工作表",
            )
            retry_excel_call(lambda: worksheet.Activate(), "激活工作统计工作表")
            icon_file = self._find_excel_icon_file()

            try:
                for row_index, record in enumerate(records, 2):
                    if not os.path.exists(record.source_file):
                        continue

                    def add_attachment():
                        cell = worksheet.Cells(row_index, 6)
                        options = {
                            "Filename": os.path.abspath(record.source_file),
                            "Link": False,
                            "DisplayAsIcon": True,
                            "IconLabel": os.path.basename(record.source_file),
                            "Left": cell.Left + 2,
                            "Top": cell.Top + 2,
                        }
                        if icon_file:
                            options["IconFileName"] = icon_file
                            options["IconIndex"] = 0
                        return worksheet.OLEObjects().Add(**options)

                    retry_excel_call(
                        add_attachment,
                        f"插入附件图标: {os.path.basename(record.source_file)}",
                    )
            except Exception as exc:
                raise RuntimeError(
                    "Excel COM插入OLE附件图标失败。请确认本机Microsoft Excel"
                    "允许通过“插入对象”功能嵌入文件，且未被Office安全策略禁用。"
                ) from exc

            retry_excel_call(lambda: workbook.Save(), "保存工作统计Excel文件")
        finally:
            if workbook is not None:
                workbook.Close(SaveChanges=True)
            if excel is not None:
                excel.Quit()
            pythoncom.CoUninitialize()

    def _find_excel_icon_file(self) -> Optional[str]:
        """查找Excel手动插入对象时常用的图标资源文件。"""
        candidates = [
            os.path.join(
                os.environ.get("WINDIR", r"C:\Windows"),
                "Installer",
                "*",
                "xlicons.exe",
            ),
            os.path.join(
                os.environ.get("ProgramFiles", r"C:\Program Files"),
                "Microsoft Office",
                "root",
                "vfs",
                "Windows",
                "Installer",
                "*",
                "xlicons.exe",
            ),
            os.path.join(
                os.environ.get("ProgramFiles", r"C:\Program Files"),
                "Microsoft Office",
                "root",
                "Office*",
                "xlicons.exe",
            ),
            os.path.join(
                os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
                "Microsoft Office",
                "root",
                "Office*",
                "xlicons.exe",
            ),
        ]
        for pattern in candidates:
            for path in glob.glob(pattern):
                if os.path.isfile(path):
                    return path
        return None

    def _show_records(self, records: list[WorkStatRecord]):
        self.results_tree.delete(*self.results_tree.get_children())
        for record in records:
            self.results_tree.insert(
                "",
                tk.END,
                values=(
                    record.sequence,
                    record.prefix,
                    record.title,
                    record.applicant_unit,
                    record.applicant,
                    record.status,
                ),
            )

    def _reset_ui_state(self):
        self.is_processing = False
        self._set_controls_state(tk.NORMAL)
        self.status_var.set("就绪")

    def set_initial_focus(self):
        try:
            self.import_files_btn.focus_set()
        except Exception:
            pass
