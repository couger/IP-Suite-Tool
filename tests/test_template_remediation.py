from pathlib import Path
import shutil
import sys

from openpyxl import Workbook

from ip_work_stats import WorkStats
from utils.table_schema import APPLICATION_COLUMNS, PLAN_COLUMNS

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from check_public_content import inspect_xlsx, inspect_tree


def test_customer_layout_cannot_be_reintroduced_as_a_renamed_example(tmp_path):
    path = tmp_path / 'requests.xlsx'
    book = Workbook()
    book.active.append([f'Unreviewed column {index}' for index in range(22)])
    book.save(path)
    assert 'example schema differs from the reviewed generic schema' in inspect_xlsx(path)


def test_hidden_extra_sheet_in_a_generic_example_is_blocked(tmp_path):
    path = tmp_path / 'requests.xlsx'
    book = Workbook()
    book.active.append(APPLICATION_COLUMNS)
    sheet = book.create_sheet('unreviewed')
    sheet.sheet_state = 'hidden'
    book.save(path)
    assert 'example must contain only the reviewed generic sheet' in inspect_xlsx(path)


def test_quarantined_source_cannot_return_to_publication_tree(tmp_path):
    (tmp_path / 'legacy').mkdir()
    (tmp_path / 'legacy/old.py').write_text('print("synthetic")', encoding='utf-8')
    assert any('quarantined' in issue for issue in inspect_tree(tmp_path))


def test_generic_schema_is_minimal_and_general_file_names_work():
    assert len(APPLICATION_COLUMNS) == 10
    assert len(PLAN_COLUMNS) == 4
    stats = object.__new__(WorkStats)
    record = stats._read_record(1, str(ROOT / 'examples/requests.xlsx'))
    assert record.status == '成功'
    assert record.prefix == ''
    assert record.title == 'requests'
    assert record.applicant_unit == '示例组织'
    assert record.applicant == '示例联系人'


def test_even_generic_layout_rejects_document_author_metadata(tmp_path):
    path = tmp_path / 'requests.xlsx'
    book = Workbook()
    book.active.append(APPLICATION_COLUMNS)
    book.properties.creator = 'Synthetic test author'
    book.save(path)
    assert 'example document metadata not allowed' in inspect_xlsx(path)


def test_real_values_cannot_replace_reviewed_examples_silently(tmp_path):
    (tmp_path / 'src').mkdir()
    shutil.copytree(ROOT / 'examples', tmp_path / 'examples')
    from openpyxl import load_workbook
    path = tmp_path / 'examples/requests.xlsx'
    book = load_workbook(path)
    book.active['I2'] = 'Changed synthetic organization'
    book.save(path)
    book.close()
    assert any('content changed since synthetic-example review' in issue
               for issue in inspect_tree(tmp_path))
