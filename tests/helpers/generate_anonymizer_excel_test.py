#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""生成IP脱敏Excel测试文件并运行脱敏验证。"""

from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SRC_DIR = PROJECT_ROOT / "src"
OUTPUT_DIR = PROJECT_ROOT / "test_outputs"
INPUT_XLSX = OUTPUT_DIR / "anonymizer_ip_cases.xlsx"
OUTPUT_XLSX = OUTPUT_DIR / "anonymizer_ip_cases_anonymized.xlsx"

sys.path.insert(0, str(SRC_DIR))

from models.model_anonymizer_core import IPAnonymizerCore  # noqa: E402
from models.model_anonymizer_processor import FileProcessor  # noqa: E402


CASES = [
    ("IPv4 private", "Login from 192.168.1.1"),
    ("IPv4 public", "DNS 8.8.8.8 and 1.1.1.1"),
    ("IPv4 adjacent text", "prefixA10.0.0.1suffix"),
    ("IPv4 invalid should stay", "Invalid 999.999.999.999 is not an IP"),
    ("IPv6 loopback", "loopback::1done"),
    ("IPv6 compressed", "route 2001:db8::1 ok"),
    ("IPv6 full", "addr=2001:0db8:85a3:0000:0000:8a2e:0370:7334"),
    ("IPv6 link-local", "gateway fe80::1ff:fe23:4567:890a"),
    ("IPv6 FC00 full", "src fc00:0:0:0:0:0:0:1"),
    ("IPv6 FC00 compressed", "dst=fc00::1"),
    ("IPv4-mapped dotted", "::ffff:192.168.1.1"),
    ("IPv4-mapped hex", "::ffff:c0a8:0101"),
    ("IPv4-compatible", "::192.168.1.1"),
    ("IPv6 embedded prefix", "2001:db8::192.168.1.1"),
    ("FC00 embedded IPv4", "fc00::192.168.1.1"),
    (
        "FC00 range",
        "range fc00:0:0:0:0:0:0:1-fc00:0:0:0:0:0:0:2",
    ),
    ("Chinese paragraph", "连接到::ffff:10.0.0.1服务器，再访问2001:db8::2。"),
    ("Multiple in one cell", "A=10.1.2.3;B=2001:db8:1::10;C=fc00::abcd"),
]


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = "A2"
    widths = [24, 72, 18, 72, 16]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def create_input_workbook():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    core = IPAnonymizerCore(partial_anonymize=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "IP脱敏用例"
    ws.append(["类型", "原始内容", "预期识别数量", "核心直接脱敏结果", "处理后剩余IP数"])

    for case_type, original in CASES:
        processed, stats = core.process_text(original)
        remaining = core.count_ips_in_text(processed).total_count
        ws.append([case_type, original, stats.total_count, processed, remaining])

    style_sheet(ws)
    wb.save(INPUT_XLSX)


def run_file_processor():
    processor = FileProcessor(partial_anonymize=True)
    result = processor.process_xlsx_file(str(INPUT_XLSX), str(OUTPUT_XLSX))
    if result.status != "success":
        raise RuntimeError(result.error_message)
    return result


def verify_output_workbook():
    core = IPAnonymizerCore(partial_anonymize=True)
    wb = load_workbook(OUTPUT_XLSX, data_only=True)
    failures = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    count = core.count_ips_in_text(value).total_count
                    if count:
                        failures.append((ws.title, value, count))
    wb.close()
    if failures:
        details = "\n".join(f"{sheet}: {value} ({count})" for sheet, value, count in failures)
        raise AssertionError(f"脱敏后仍发现有效IP:\n{details}")


def main():
    create_input_workbook()
    result = run_file_processor()
    verify_output_workbook()
    print(f"input={INPUT_XLSX}")
    print(f"output={OUTPUT_XLSX}")
    print(f"status={result.status}")
    print(f"ip_count={result.ip_count}")


if __name__ == "__main__":
    main()
