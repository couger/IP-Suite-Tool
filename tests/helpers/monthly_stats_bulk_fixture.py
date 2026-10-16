#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


BASE = Path(__file__).resolve().parents[2] / "test_outputs"
BATCH = BASE / "\u81ea\u52a8\u5316\u6279\u91cf\u6d4b\u8bd5_20260905"


UNITS = [f"示例单位{index}" for index in range(1, 9)]
PEOPLE = [f"示例人员{index}" for index in range(1, 9)]
SCENES = [f"示例业务{index}" for index in range(1, 9)]


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="D9EAD3")
    border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=5):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for idx, width in enumerate([22, 34, 22, 36, 26], 1):
        ws.column_dimensions[chr(64 + idx)].width = width


def make_workbook(path: Path, index: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "\u7533\u8bf7\u8868"
    ws.append(
        [
            "\u4e1a\u52a1\u8bf4\u660e",
            "组织",
            "联系人",
            "\u5730\u5740\u8d44\u6e90\u660e\u7ec6",
            "\u5176\u4ed6\u8bf4\u660e",
        ]
    )
    ws.append(
        [
            f"\u7b2c{index}\u6279\u7533\u8bf7",
            UNITS[index % len(UNITS)],
            PEOPLE[index % len(PEOPLE)],
            "\u521d\u59cb\u884c",
            "",
        ]
    )
    ws.append(
        [
            "\u4e2d\u95f4\u7a7a\u884c\u6d4b\u8bd5",
            "",
            "",
            "",
            "\u7528\u4e8e\u9a8c\u8bc1\u6700\u540e\u975e\u7a7a\u5355\u5143\u683c",
        ]
    )
    ws.append(
        [
            "\u8865\u5145\u7533\u8bf7",
            UNITS[(index + 2) % len(UNITS)] + "-\u65e7\u503c",
            PEOPLE[(index + 2) % len(PEOPLE)] + "-\u65e7\u503c",
            "",
            "",
        ]
    )
    ws.append(
        [
            "\u6700\u7ec8\u786e\u8ba4",
            UNITS[(index + 4) % len(UNITS)] + f"-\u6700\u7ec8{index:02d}",
            PEOPLE[(index + 4) % len(PEOPLE)] + f"-\u6700\u7ec8{index:02d}",
            "",
            "\u6700\u540e\u4e00\u884c\u5e94\u88ab\u7edf\u8ba1",
        ]
    )
    # Work statistics reads the last non-empty applicant fields of the first sheet.
    # Keep IP samples on a separate sheet so they cannot become applicant data.
    style_sheet(ws)
    ws = wb.create_sheet("IP样本")
    ws.append(["\u4ee5\u4e0b\u4e3a\u53ef\u8131\u654fIP\u6837\u672c", "", "", "", ""])
    rows = [
        ["\u6e90\u5730\u5740", "\u76ee\u7684\u5730\u5740", "\u5907\u6ce8"],
        ["\u670d\u52a1\u5668192.168.10.11\u5f85\u8131\u654f", "2001:db8::10", "IPv4\u4e0eIPv6\u6df7\u5408"],
        ["10.20.30.40", "fc00:0:0:0:0:0:0:1", "fc00\u5730\u5740"],
        ["::ffff:172.16.5.9", "2001:db8::192.168.1.1", "IPv4-mapped\u548c\u5d4c\u5165IPv4"],
        ["\u8bbf\u95ee8.8.8.8\u548c1.1.1.1", "fe80::1", "\u516c\u5171DNS"],
        ["\u4e2d\u6587\u76f8\u90bbIP", f"\u8bbe\u5907A10.{index % 255}.{(index * 3) % 255}.1\u7aef\u53e3", f"\u94fe\u8def2001:db8::{index}\u5df2\u5f00\u901a"],
    ]
    for row in rows:
        ws.append(row + [f"\u6837\u672c{index}"])

    style_sheet(ws)
    wb.save(path)


def main() -> None:
    BATCH.mkdir(parents=True, exist_ok=True)
    created = []
    for index in range(1, 49):
        path = BATCH / f"request-{index:04d}-{SCENES[index % len(SCENES)]}.xlsx"
        make_workbook(path, index)
        created.append(path)

    manifest = BATCH / "\u751f\u6210\u6e05\u5355.txt"
    manifest.write_text("\n".join(str(path) for path in created), encoding="utf-8")
    print(f"created={len(created)}")
    print(f"batch={BATCH}")
    print(f"manifest={manifest}")


if __name__ == "__main__":
    main()
