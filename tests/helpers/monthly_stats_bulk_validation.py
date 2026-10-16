#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import json
import os
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "src"
BATCH = ROOT / "test_outputs" / "\u81ea\u52a8\u5316\u6279\u91cf\u6d4b\u8bd5_20260905"

sys.path.insert(0, str(SRC))

from ip_work_stats import WorkStats  # noqa: E402
from models.model_anonymizer_processor import FileProcessor  # noqa: E402


def load_target_files() -> list[str]:
    manifest = BATCH / "\u751f\u6210\u6e05\u5355.txt"
    if manifest.exists():
        files = [
            line.strip()
            for line in manifest.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
    else:
        files = sorted(str(path) for path in BATCH.glob("*.xlsx"))
    if not files:
        raise RuntimeError(f"No .xlsx files found in {BATCH}")
    return files


def validate_work_stats(files: list[str]) -> dict:
    stats = object.__new__(WorkStats)
    records = [WorkStats._read_record(stats, index, path) for index, path in enumerate(files, 1)]

    failed = [record for record in records if record.status != "\u6210\u529f"]
    missing_filename = [record.source_file for record in records if not (record.prefix or record.title)]
    missing_people = [
        record.source_file
        for record in records
        if not record.applicant_unit or not record.applicant
    ]

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = BATCH / f"\u5de5\u4f5c\u7edf\u8ba1_\u6279\u91cf\u9a8c\u8bc1_{timestamp}.xlsx"
    WorkStats._save_records(stats, records, str(output_path), embed_attachments=True)

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        row_count = sheet.max_row - 1
        source_paths = [
            sheet.cell(row=row_index, column=7).value
            for row_index in range(2, sheet.max_row + 1)
        ]
    finally:
        workbook.close()

    attachment_parts = 0
    try:
        import zipfile

        with zipfile.ZipFile(output_path) as archive:
            attachment_parts = len(
                [
                    name
                    for name in archive.namelist()
                    if name.startswith("xl/embeddings/")
                ]
            )
    except Exception:
        attachment_parts = -1

    return {
        "records": len(records),
        "export_rows": row_count,
        "failed_records": len(failed),
        "missing_filename": len(missing_filename),
        "missing_applicant_info": len(missing_people),
        "source_path_rows": len([path for path in source_paths if path]),
        "attachment_parts": attachment_parts,
        "output_path": str(output_path),
    }


def workbook_has_detectable_ip(processor: FileProcessor, file_path: str) -> bool:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if isinstance(value, str):
                        if processor.ip_core.count_ips_in_text(value).total_count:
                            return True
        return False
    finally:
        workbook.close()


def validate_anonymizer(files: list[str]) -> dict:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = BATCH / f"\u8131\u654f\u6279\u91cf\u9a8c\u8bc1_{timestamp}"
    processor = FileProcessor(partial_anonymize=True)
    results, stats = processor.process_files_batch_with_structure(files, str(output_dir))

    residual_files = []
    for result in results:
        if result.status == "success" and result.output_path:
            if workbook_has_detectable_ip(processor, result.output_path):
                residual_files.append(result.output_path)

    return {
        "input_files": len(files),
        "successful_files": stats.successful_files,
        "failed_files": stats.failed_files,
        "skipped_files": stats.skipped_files,
        "total_ips_found": stats.total_ips_found,
        "residual_files": residual_files,
        "output_dir": str(output_dir),
    }


def main() -> None:
    files = load_target_files()
    work_stats = validate_work_stats(files)
    anonymizer = validate_anonymizer(files)
    passed = (
        work_stats["records"] == len(files)
        and work_stats["export_rows"] == len(files)
        and work_stats["failed_records"] == 0
        and work_stats["missing_filename"] == 0
        and work_stats["missing_applicant_info"] == 0
        and work_stats["source_path_rows"] == len(files)
        and work_stats["attachment_parts"] >= len(files)
        and anonymizer["successful_files"] == len(files)
        and anonymizer["failed_files"] == 0
        and anonymizer["residual_files"] == []
    )
    report = {
        "passed": passed,
        "batch": str(BATCH),
        "work_stats": work_stats,
        "anonymizer": anonymizer,
    }
    report_path = BATCH / f"\u81ea\u52a8\u5316\u9a8c\u8bc1\u62a5\u544a_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    print(f"report_path={report_path}")
    if not passed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
