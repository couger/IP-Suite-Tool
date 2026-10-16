#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import os
import sys

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(__file__)), "src"))

from ip_work_stats import WorkStats
from ip_anonymizer import IPAnonymizer
from models.model_anonymizer_processor import FileProcessor, ProcessingResult, ProcessingStats


def _make_workbook(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["其他列", "申请单位名称", "申请人信息"])
    sheet.append(["a", "第一单位", "第一人"])
    sheet.append(["b", "", ""])
    sheet.append(["c", "最终单位", "最终申请人"])
    workbook.save(path)


def _make_workbook_with_offset_header(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["说明", "本行不是标题", ""])
    sheet.append(["其他列", "申请单位名称", "申请人信息"])
    sheet.append(["a", "第一单位", "第一人"])
    sheet.append(["b", "最终单位", "最终申请人"])
    workbook.save(path)


def test_work_stats_extracts_prefix_title_and_last_non_empty_values(tmp_path):
    file_path = tmp_path / "pl202601063001-测试申请表.xlsx"
    _make_workbook(file_path)

    work_stats = object.__new__(WorkStats)
    record = WorkStats._read_record(work_stats, 1, str(file_path))

    assert record.sequence == 1
    assert record.prefix == "pl202601063001"
    assert record.title == "测试申请表"
    assert record.applicant_unit == "最终单位"
    assert record.applicant == "最终申请人"
    assert record.status == "成功"


def test_work_stats_finds_applicant_header_below_first_row(tmp_path):
    file_path = tmp_path / "pl202601063001-多行标题申请表.xlsx"
    _make_workbook_with_offset_header(file_path)

    work_stats = object.__new__(WorkStats)
    record = WorkStats._read_record(work_stats, 1, str(file_path))

    assert record.status == "成功"
    assert record.applicant_unit == "最终单位"
    assert record.applicant == "最终申请人"


def test_work_stats_allows_empty_remaining_filename_and_exports_columns(tmp_path):
    file_path = tmp_path / "add202604013002.xlsx"
    output_path = tmp_path / "工作统计.xlsx"
    _make_workbook(file_path)

    work_stats = object.__new__(WorkStats)
    record = WorkStats._read_record(work_stats, 1, str(file_path))
    WorkStats._save_records(
        work_stats, [record], str(output_path), embed_attachments=False
    )

    workbook = load_workbook(output_path, data_only=True)
    sheet = workbook.active

    assert [sheet.cell(1, col).value for col in range(1, 6)] == [
        "序号",
        "前缀",
        "文件名剩余部分",
        "组织",
        "联系人",
    ]
    assert sheet.cell(1, 6).value == "附件"
    assert sheet.cell(2, 2).value == "add202604013002"
    assert sheet.cell(2, 3).value is None
    assert sheet.cell(2, 6).value is None
    assert sheet.cell(2, 7).value == str(file_path)

    workbook.close()


def test_anonymizer_work_stats_input_includes_failed_original_files(tmp_path):
    source_file = tmp_path / "pl202601063001-失败原文件.xlsx"
    output_file = tmp_path / "pl202601063002-成功输出.xlsx"
    _make_workbook(source_file)
    _make_workbook(output_file)

    results = [
        ProcessingResult(
            file_path=str(source_file),
            status="error",
            original_size=1,
            processed_size=0,
            ip_count=0,
            error_message="测试失败",
        ),
        ProcessingResult(
            file_path=str(source_file),
            status="success",
            original_size=1,
            processed_size=1,
            ip_count=1,
            output_path=str(output_file),
        ),
    ]

    anonymizer = object.__new__(IPAnonymizer)

    assert IPAnonymizer._get_work_stats_input_files(anonymizer, results) == [
        str(source_file),
        str(output_file),
    ]


def test_work_stats_includes_txt_and_word_without_applicant_lookup(tmp_path):
    txt_file = tmp_path / "add202601063003-文本申请.txt"
    docx_file = tmp_path / "pl202601063004-Word申请.docx"
    ignored_file = tmp_path / "pl202601063005-忽略.pdf"
    txt_file.write_text("服务器 192.168.1.1", encoding="utf-8")
    docx_file.write_text("placeholder", encoding="utf-8")
    ignored_file.write_text("placeholder", encoding="utf-8")

    work_stats = object.__new__(WorkStats)

    assert WorkStats._is_supported_file(work_stats, str(txt_file))
    assert WorkStats._is_supported_file(work_stats, str(docx_file))
    assert not WorkStats._is_supported_file(work_stats, str(ignored_file))

    txt_record = WorkStats._read_record(work_stats, 1, str(txt_file))
    docx_record = WorkStats._read_record(work_stats, 2, str(docx_file))

    assert txt_record.status == "成功"
    assert txt_record.applicant_unit == ""
    assert txt_record.applicant == ""
    assert "不读取组织和联系人" in txt_record.message
    assert docx_record.status == "成功"
    assert docx_record.applicant_unit == ""
    assert docx_record.applicant == ""
    assert "不读取组织和联系人" in docx_record.message


def test_processing_stats_uses_instance_unprocessed_file_lists():
    assert ProcessingStats().unprocessed_files is not ProcessingStats().unprocessed_files


def test_file_processor_validates_binary_file_access_and_docx_output_path(tmp_path):
    from docx import Document

    docx_file = tmp_path / "source.docx"
    output_file = tmp_path / "output.docx"
    document = Document()
    document.add_paragraph("服务器地址 192.168.1.1")
    document.save(docx_file)

    processor = FileProcessor(partial_anonymize=True)

    assert processor.validate_file_access(str(docx_file))

    result = processor.process_docx_file(str(docx_file), str(output_file))

    assert result.status == "success"
    assert result.output_path == str(output_file)
    assert output_file.exists()


def test_docx_anonymization_preserves_run_formatting(tmp_path):
    from docx import Document

    docx_file = tmp_path / "formatted.docx"
    output_file = tmp_path / "formatted_output.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("服务器 ")
    ip_run = paragraph.add_run("192.168.1.1")
    ip_run.bold = True
    paragraph.add_run(" 结束")
    document.save(docx_file)

    processor = FileProcessor(partial_anonymize=True)
    result = processor.process_docx_file(str(docx_file), str(output_file))

    assert result.status == "success"

    output_doc = Document(output_file)
    output_paragraph = output_doc.paragraphs[0]

    assert output_paragraph.runs[0].text == "服务器 "
    assert output_paragraph.runs[1].text == "192.x.x.x"
    assert output_paragraph.runs[1].bold is True
    assert output_paragraph.runs[2].text == " 结束"
