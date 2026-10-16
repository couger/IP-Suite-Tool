import io
import logging
from pathlib import Path
import shutil
import zipfile

from openpyxl import Workbook, load_workbook

from ip_work_stats import WorkStats
from models.model_ip_allocator_core import IPAllocationCore
from utils.app_paths import config_path, resource_path
from utils.logger_config import PrivateDataSafeFormatter
from utils.table_schema import read_display_table, canonical_column

ROOT = Path(__file__).resolve().parents[1]
EXAMPLES = ROOT / "examples"


def test_published_templates_support_loading_allocation_and_statistics(tmp_path):
    app = EXAMPLES / "requests.xlsx"
    master = EXAMPLES / "allocations.xlsx"
    core = IPAllocationCore()
    assert core.load_application_table(str(app))
    assert core.load_master_table(str(master))
    assert core.load_plan_table(str(EXAMPLES / "network_plan.xlsx"))
    result = core.process_allocation("DEMO-VALIDATION")
    assert result["allocated"] == 1
    assert result["skipped"] == 1
    assert result["error"] == 0
    assert core.save_application_table(str(tmp_path / "allocated.xlsx"))
    assert core.save_master_table(str(tmp_path / "master.xlsx"))
    assert len(load_workbook(tmp_path / "master.xlsx").active["A"]) == 3
    stats = object.__new__(WorkStats)
    record = stats._read_record(1, str(app))
    assert record.prefix == ""
    assert record.applicant_unit == "示例组织"
    assert record.applicant == "示例联系人"
    stats._save_records([record], str(tmp_path / "statistics.xlsx"), embed_attachments=False)
    assert load_workbook(tmp_path / "statistics.xlsx").active.max_row == 2


def test_site_values_sheet_names_and_star_spacing_are_not_coupled(tmp_path):
    source = EXAMPLES / "requests.xlsx"
    target = tmp_path / "renamed.xlsx"
    shutil.copyfile(source, target)
    book = load_workbook(target)
    sheet = book.active
    sheet.title = "Independent sheet"
    sheet["B1"] = " * 站点 * "
    sheet["C1"] = "设备点位 ＊"
    sheet["D1"] = " 所属业务 * "
    sheet["B2"] = "任意测试位置"
    sheet["C2"] = "任意测试资源"
    book.save(target)
    core = IPAllocationCore()
    assert core.load_application_table(str(target))
    assert core.application_df.iloc[0]["位置"] == "任意测试位置"
    columns, records = read_display_table(target)
    assert columns[1] == "* 站点 *"
    assert len(records) == 2  # Pending records have a blank application ID.
    assert canonical_column("IPv4地址(按需)") == "IPv4地址"
    output = tmp_path / "legacy-roundtrip.xlsx"
    assert core.save_application_table(str(output))
    assert load_workbook(output).active["B1"].value == " * 站点 * "


def test_old_column_aliases_do_not_require_or_export_a_customer_profile(tmp_path):
    path = tmp_path / "legacy-import.xlsx"
    book = Workbook()
    book.active.append(["站点", "设备点位", "所属业务*", "申请单位", "申请人", "自定义说明"])
    book.active.append(["测试位置", "测试资源", "演示", "测试组织", "测试联系人", "本地保留"])
    book.save(path)
    core = IPAllocationCore()
    assert core.load_application_table(str(path))
    assert core.load_master_table(str(tmp_path / "missing-master.xlsx"))
    assert core.process_allocation("DEMO")['allocated'] == 1
    output = tmp_path / "legacy-export.xlsx"
    assert core.save_application_table(str(output))
    restored = load_workbook(output).active
    assert restored["A1"].value == "站点"
    assert restored["F2"].value == "本地保留"


def test_colliding_aliases_are_rejected(tmp_path):
    path = tmp_path / "ambiguous.xlsx"
    book = Workbook()
    book.active.append(["位置", "站点", "资源标识", "业务类别"])
    book.active.append(["A", "B", "asset-a", "演示"])
    book.save(path)
    assert not IPAllocationCore().load_application_table(str(path))


def test_runtime_state_and_frozen_resources_are_separate(tmp_path, monkeypatch):
    monkeypatch.setenv("IP_SUITE_DATA_DIR", str(tmp_path / "private-state"))
    path = config_path("prefixes.json")
    assert path.parent == tmp_path / "private-state/cfgs"
    assert not path.is_relative_to(ROOT)
    monkeypatch.setattr("sys._MEIPASS", str(tmp_path / "unpacked"), raising=False)
    assert resource_path("defaults", "prefixes.json") == tmp_path / "unpacked/defaults/prefixes.json"


def test_logs_do_not_include_message_arguments_fstrings_or_exception_data():
    secret = "SYNTHETIC-PRIVATE-PAYLOAD"
    formatter = PrivateDataSafeFormatter()
    for message, args in [("Value %s", (secret,)), (f"Value {secret}", ())]:
        record = logging.LogRecord("test", logging.ERROR, __file__, 1, message, args, None)
        assert secret not in formatter.format(record)
    try:
        raise ValueError(secret)
    except ValueError:
        import sys
        record = logging.LogRecord("test", logging.ERROR, __file__, 1, secret, (), sys.exc_info())
        formatted = formatter.format(record)
        assert secret not in formatted
        assert "ValueError" in formatted


def test_rebuilt_examples_have_no_hidden_links_or_legacy_properties():
    import sys
    sys.path.insert(0, str(ROOT / "scripts"))
    from check_public_content import inspect_xlsx
    for path in EXAMPLES.glob("*.xlsx"):
        assert inspect_xlsx(path) == []


def test_privacy_check_rejects_forced_tracking_of_runtime_data(tmp_path):
    import subprocess
    import sys
    sys.path.insert(0, str(ROOT / "scripts"))
    from check_public_content import inspect_tree
    subprocess.run(["git", "init", "-q", str(tmp_path)], check=True)
    (tmp_path / "runtime").mkdir()
    (tmp_path / "runtime/private.log").write_text("SYNTHETIC DATA", encoding="utf-8")
    subprocess.run(["git", "add", "runtime/private.log"], cwd=tmp_path, check=True)
    assert "runtime/private.log: runtime/generated directory is tracked" in inspect_tree(tmp_path)
