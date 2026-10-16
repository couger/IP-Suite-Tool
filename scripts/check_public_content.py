"""Fail closed on runtime files, private paths, or unsafe example packages.

This check supplements review and a credential scanner; it cannot identify every
company-specific value or determine ownership of source code.
"""
import argparse
import hashlib
import json
import pathlib
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile

ROOT = pathlib.Path(__file__).resolve().parents[1]
ALLOWED_XLSX = {
    "examples/network_plan.xlsx",
    "examples/requests.xlsx",
    "examples/allocations.xlsx",
}
EXAMPLE_SCHEMAS = {
    "network_plan.xlsx": ("位置", "IPv4网段", "IPv6网段", "用途"),
    "requests.xlsx": ("记录编号", "位置", "资源标识", "业务类别", "IPv4地址", "IPv6地址", "网关", "掩码", "组织", "联系人"),
    "allocations.xlsx": ("记录编号", "位置", "资源标识", "业务类别", "IPv4地址", "IPv6地址", "网关", "掩码", "组织", "联系人"),
}
ALLOWED_CONFIG = {"src/cfgs/cfg_mapping_config.json"}
EXCLUDED = {".git", ".venv", "venv", "__pycache__", ".pytest_cache", "build", "build-clean", "dist", "bin", "node_modules", "test_outputs", "data", "runtime"}
PRIVATE_PATH = re.compile(r"(?i)[a-z]:[/\\](?:users|dev|工作)[/\\]")
EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@([A-Za-z0-9.-]+\.[A-Za-z]{2,})")


def text_issues(value):
    issues = []
    if PRIVATE_PATH.search(value):
        issues.append("personal or business absolute path")
    if ("wx" + "id_") in value or ("WeChat" + " Files") in value:
        issues.append("personal messaging storage identifier")
    if any(m.group(1).lower() not in {"example.com", "example.org", "example.net", "users.noreply.github.com"} for m in EMAIL.finditer(value)):
        issues.append("non-example email")
    return issues


def inspect_xlsx(path):
    issues = []
    try:
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                lower = name.lower()
                if path.name in EXAMPLE_SCHEMAS and lower.startswith("docprops/"):
                    issues.append("example document metadata not allowed")
                if any(part in lower for part in ("externallink", "embedding", "vbaproject", "customxml", "customstorage", "comments", "docprops/custom")):
                    issues.append("unexpected external/hidden package component")
                if name.endswith((".xml", ".rels")):
                    element = ET.fromstring(archive.read(name))
                    content = " ".join(element.itertext()) + " ".join(str(v) for e in element.iter() for v in e.attrib.values())
                    issues.extend(text_issues(content))
                    if any(e.attrib.get("TargetMode") == "External" for e in element.iter()):
                        issues.append("external relationship")
    except (OSError, zipfile.BadZipFile, ET.ParseError):
        issues.append("invalid workbook package")
    if path.name in EXAMPLE_SCHEMAS:
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                if len(workbook.worksheets) != 1:
                    issues.append("example must contain only the reviewed generic sheet")
                sheet = workbook.worksheets[0]
                headers = tuple(cell.value for cell in next(sheet.iter_rows(max_row=1)) if cell.value is not None)
                if headers != EXAMPLE_SCHEMAS[path.name]:
                    issues.append("example schema differs from the reviewed generic schema")
                if any(cell.data_type == 'f' for row in sheet for cell in row):
                    issues.append("example formula not allowed")
                if len(workbook.defined_names):
                    issues.append("example named range not allowed")
            finally:
                workbook.close()
        except Exception:
            issues.append("example schema could not be validated")
    return sorted(set(issues))


def inspect_tree(root=ROOT):
    findings = []
    if (root / "src").is_dir():
        manifest_path = root / "examples/manifest.json"
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            if set(manifest) != {pathlib.Path(name).name for name in ALLOWED_XLSX}:
                raise ValueError("Unexpected example manifest")
            for filename, expected_hash in manifest.items():
                example = root / "examples" / filename
                if hashlib.sha256(example.read_bytes()).hexdigest() != expected_hash:
                    findings.append(f"examples/{filename}: content changed since synthetic-example review")
        except (OSError, ValueError, TypeError):
            findings.append("reviewed synthetic example manifest is missing or invalid")
    tracked = set()
    if (root / ".git").exists():
        result = subprocess.run(["git", "ls-files", "-z"], cwd=root, capture_output=True, check=True)
        tracked = {name for name in result.stdout.decode("utf-8").split("\0") if name}
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        relative = path.relative_to(root)
        name = relative.as_posix()
        if any(part in EXCLUDED for part in relative.parts):
            if name in tracked:
                findings.append(f"{name}: runtime/generated directory is tracked")
            continue
        issues = []
        if relative.parts[0] == "legacy":
            issues.append("unreviewed historical source is quarantined outside the publication tree")
        if path.suffix.lower() == ".xlsx":
            if name not in ALLOWED_XLSX:
                issues.append("workbook outside approved examples")
            issues.extend(inspect_xlsx(path))
        elif path.suffix.lower() in {".xls", ".xlsm", ".docx", ".bundle", ".patch", ".exe", ".dll", ".log", ".zip", ".7z"}:
            issues.append("private data or generated binary in source tree")
        else:
            if name.startswith("src/cfgs/") and name not in ALLOWED_CONFIG:
                issues.append("runtime configuration in source resources")
            try:
                content = path.read_text(encoding="utf-8-sig")
            except (UnicodeError, OSError):
                content = ""
            issues.extend(text_issues(content))
        findings.extend(f"{name}: {issue}" for issue in set(issues))
    return sorted(findings)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--binary", type=pathlib.Path, help="Also inspect a PyInstaller executable's packaged data")
    args = parser.parse_args()
    issues = inspect_tree()
    if args.binary:
        import importlib.resources
        from PyInstaller.archive.readers import CArchiveReader
        reader = CArchiveReader(str(args.binary))
        allowed = {"cfgs/cfg_mapping_config.json", "defaults/cfg_ipv6_instance_prefix.json"}
        for name in reader.toc:
            normalized = name.replace("\\", "/")
            # python-docx needs its own blank template. It is dependency data,
            # not a user document; require byte identity with the installed package.
            if normalized == "docx/templates/default.docx":
                template = importlib.resources.files("docx").joinpath("templates/default.docx").read_bytes()
                if reader.extract(name) != template:
                    issues.append("binary: modified dependency Word template")
                continue
            if normalized.startswith(("logs/", "examples/")) or normalized.lower().endswith((".xlsx", ".xls", ".docx", ".log")):
                issues.append(f"binary: forbidden data entry {normalized}")
            if normalized.startswith(("cfgs/", "defaults/")):
                if normalized not in allowed:
                    issues.append(f"binary: unapproved configuration {normalized}")
                else:
                    issues.extend(f"binary {normalized}: {issue}" for issue in text_issues(reader.extract(name).decode("utf-8")))
    for issue in issues:
        print(issue)
    print(f"Public content check: {'FAILED' if issues else 'passed'}")
    return 1 if issues else 0


if __name__ == "__main__":
    sys.exit(main())
